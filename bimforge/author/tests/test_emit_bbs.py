import pytest


def test_emit_bbs_writes_schedule(tmp_path):
    import uuid
    pytest.importorskip("openpyxl")           # optional, non-bundled dependency
    from openpyxl import load_workbook
    from bimforge.author.emit_bbs import emit_bbs
    from bimforge.author.model import Element, StructureModel
    from bimforge.author.project import ProjectContext
    m = StructureModel("t", ProjectContext(name="t", guid_ns=uuid.UUID(int=1)))
    m.add(Element("WWE #5 reinforcing", "IfcReinforcingBar",
                  quantities={"count": 21, "length_ft": 212.6, "weight_lb": 221.7},
                  description="#5 k"))
    m.add(Element("WWE #6 reinforcing", "IfcReinforcingBar",
                  quantities={"count": 2, "length_ft": 38.0, "weight_lb": 57.1},
                  description="#6 L"))
    emit_bbs(m, str(tmp_path / "ww.bbs.xlsx"))
    wb = load_workbook(tmp_path / "ww.bbs.xlsx")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Mark", "Size", "Count", "Length (ft)", "Weight (lb)")
    assert any(r[0] == "WWE #5 reinforcing" and r[4] == 221.7 for r in rows)
    assert rows[-1][0] == "TOTAL" and rows[-1][4] == pytest.approx(221.7 + 57.1)
