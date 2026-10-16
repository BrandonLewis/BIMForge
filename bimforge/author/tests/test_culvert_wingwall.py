import uuid
import pytest
from bimforge.author.gate import run_pre
from bimforge.author.project import ProjectContext
from bimforge.author.families.culvert_wingwall import build

CTX = ProjectContext(name="ww", guid_ns=uuid.UUID(int=22))
STRAIGHT = {"wingwalls": [
    {"name": "WW1", "kind": "straight", "wall_h_ft": 8.0, "length_ft": 18.0,
     "min_h_ft": 1.0, "stem_t_ft": 0.8333, "toe_ft": 1.0, "heel_ft": 4.0,
     "foot_t_ft": 1.0, "reinf_lb": 640}]}


def test_missing_params_raises():
    with pytest.raises(ValueError, match="missing"):
        build({}, CTX)


def test_straight_stem_and_footing_volumes():
    m = build(STRAIGHT, CTX)
    stem = m.element("WW1 wingwall stem")
    foot = m.element("WW1 wingwall footing")
    T, L, W, F = 0.8333, 18.0, 5.0, 1.0
    h_max, h_min = 8.5, 1.0                       # wall_h+0.5, min_h
    assert stem.quantities["volume_cf"] == pytest.approx(
        T * L * (h_max + h_min) / 2.0, abs=0.1)
    assert foot.quantities["volume_cf"] == pytest.approx(L * W * F, abs=0.05)
    assert run_pre(m).passed


def test_straight_reinforcing_explicit():
    m = build(STRAIGHT, CTX)
    steel = m.element("WW1 wingwall reinforcing")
    assert steel.ifc_class == "IfcReinforcingBar"
    assert steel.quantities["weight_lb"] == pytest.approx(640.0)
    assert run_pre(m).passed


STEPPED = {"tables": {"steel": {"type_e": {
    "10": {"k_bar": "#5@12", "L_bar": "#6", "L_each_wall": 2}}}},
           "wingwalls": [
    {"name": "WWE", "kind": "stepped", "stem_t_ft": 0.8333, "toe_ft": 1.0,
     "heel_ft": 4.0, "foot_t_ft": 1.0,
     "steps": [{"length_ft": 6.0, "h_ft": 8.5},   # max 8.5 -> type_e "10"
               {"length_ft": 6.0, "h_ft": 5.0},
               {"length_ft": 6.0, "h_ft": 2.0}]}]}


def test_stepped_stem_steps_are_clean_boxes():
    m = build(STEPPED, CTX)
    T = 0.8333
    for i, h in enumerate((8.5, 5.0, 2.0), start=1):
        step = m.element(f"WWE wingwall stem step {i}")
        assert step.quantities["volume_cf"] == pytest.approx(6.0 * T * h, abs=0.02)
        # each step is a clean box: L x W x H reproduces the volume (gate check)
        q = step.quantities
        assert q["length_ft"] * q["width_ft"] * q["height_ft"] == pytest.approx(
            q["volume_cf"], abs=1e-6)
    # steps must be placed in sequence, not stacked at the origin
    step2 = m.element("WWE wingwall stem step 2")
    xs = [p[0] for p in step2.geometry.base]
    assert min(xs) == pytest.approx(6.0)
    assert max(xs) == pytest.approx(12.0)
    assert run_pre(m).passed


def test_stepped_total_matches_continuous_taper_within_tol():
    # the stepped stem approximates a continuous taper h_max->h_min over L=18
    m = build(STEPPED, CTX)
    T = 0.8333
    total = sum(m.element(f"WWE wingwall stem step {i}").quantities["volume_cf"]
                for i in (1, 2, 3))
    continuous = T * 18.0 * (8.5 + 2.0) / 2.0
    assert total == pytest.approx(continuous, rel=0.03)


def test_stepped_footing_and_per_mark_reinforcing():
    m = build(STEPPED, CTX)
    foot = m.element("WWE wingwall footing")
    assert foot.quantities["volume_cf"] == pytest.approx(18.0 * 5.0 * 1.0, abs=0.05)
    # per-mark generated reinforcing (k/L + minor dowels/spacers), each a
    # weight-only IfcReinforcingBar, all positive
    rebar = [e for e in m.elements if e.ifc_class == "IfcReinforcingBar"]
    assert rebar, "expected generated per-mark reinforcing elements"
    assert all(e.quantities["weight_lb"] > 0 for e in rebar)
    # k bars are #5 (from the type_e "10" row)
    assert m.element("WWE wingwall reinforcing k #5").quantities["count"] == 21
    assert run_pre(m).passed


WARPED = {"tables": {"steel": {"warped": {
    "wall": {"8": {"slopes": {"1.5:1": {"front": "#4@12", "rear": "#4@8"}}}},
    "beam": {"by_l": {"20": [{"h_max": 20, "total": "6-#7"}]}}}}},
    "wingwalls": [
    {"name": "WWW", "kind": "warped", "wall_h_ft": 8.0, "length_ft": 18.0,
     "min_h_ft": 1.0, "stem_t_ft": 0.8333, "toe_ft": 1.0, "heel_ft": 4.0,
     "foot_t_ft": 1.0, "batter_toe": 0.125, "face_slope": "1.5:1",
     "stiffening_beam": {"w_ft": 1.0, "d_ft": 1.5},
     "toe_cutoff": {"w_ft": 1.0, "d_ft": 1.5}}]}


def test_warped_stem_beam_and_cutoff():
    m = build(WARPED, CTX)
    T, L = 0.8333, 18.0
    stem = m.element("WWW wingwall stem")
    # exact shear-preserved volume (declared closed-form, batter-independent)
    assert stem.quantities["volume_cf"] == pytest.approx(
        T * L * (8.5 + 1.0) / 2.0, rel=1e-9)
    beam = m.element("WWW wingwall stiffening beam")
    assert beam.quantities["volume_cf"] == pytest.approx(L * 1.0 * 1.5, abs=0.02)
    cutoff = m.element("WWW wingwall toe cutoff")
    assert cutoff.quantities["volume_cf"] == pytest.approx(L * 1.0 * 1.5, abs=0.02)
    assert run_pre(m).passed


def test_warped_per_mark_reinforcing_front_rear_beam():
    m = build(WARPED, CTX)
    # front/rear face + beam marks each a weight-only IfcReinforcingBar
    front = m.element("WWW wingwall reinforcing front #4")
    rear = m.element("WWW wingwall reinforcing rear #4")
    beam = m.element("WWW wingwall reinforcing beam #7")
    assert beam.quantities["count"] == 6                 # TOTAL 6-#7
    assert front.quantities["weight_lb"] > 0
    assert rear.quantities["weight_lb"] > 0
    assert run_pre(m).passed


def test_warped_minimal_no_beam_no_cutoff():
    params = {"tables": WARPED["tables"], "wingwalls": [
        {"name": "WWM", "kind": "warped", "wall_h_ft": 8.0, "length_ft": 18.0,
         "min_h_ft": 1.0, "stem_t_ft": 0.8333, "toe_ft": 1.0, "heel_ft": 4.0,
         "foot_t_ft": 1.0, "batter_toe": 0.125, "face_slope": "1.5:1"}]}
    m = build(params, CTX)
    assert m.element("WWM wingwall stem").quantities["volume_cf"] > 0
    assert m.element("WWM wingwall footing").quantities["volume_cf"] > 0
    for absent in ("WWM wingwall stiffening beam", "WWM wingwall toe cutoff"):
        with pytest.raises(KeyError):
            m.element(absent)
    assert run_pre(m).passed


def test_wingwall_paycodes_classify():
    from pathlib import Path
    from bimforge.author import paycodes as pc_mod
    p = (Path(__file__).resolve().parents[1]
         / "examples" / "caltrans-culvert-wingwall.paycodes.yaml")
    pc = pc_mod.load(p)
    m = build(WARPED, CTX)
    pc_mod.apply(m, pc)
    assert m.element("WWW wingwall stem").classifications == ["STRUCT-CONC-WW"]
    assert m.element("WWW wingwall footing").classifications == ["STRUCT-CONC-WW"]
    assert m.element("WWW wingwall stiffening beam").classifications == ["STRUCT-CONC-WW"]
    assert m.element("WWW wingwall toe cutoff").classifications == ["STRUCT-CONC-WW"]
    # per-mark reinforcing classifies to STEEL -- including the "beam" mark,
    # which the steel rule (matched FIRST) must win over the concrete beam rule
    assert m.element("WWW wingwall reinforcing beam #7").classifications == ["BAR-STEEL-WW"]
    assert m.element("WWW wingwall reinforcing front #4").classifications == ["BAR-STEEL-WW"]


def test_wingwall_example_spec_builds_check_only():
    from pathlib import Path
    from bimforge.author.families.registry import build_from_spec
    spec = (Path(__file__).resolve().parents[1]
            / "examples" / "caltrans-culvert-wingwall.spec.yaml")
    assert build_from_spec(spec, check_only=True) == 0


def test_wingwall_example_full_build(tmp_path):
    from pathlib import Path
    import yaml
    from bimforge.author.families.registry import build_from_spec
    src = (Path(__file__).resolve().parents[1]
           / "examples" / "caltrans-culvert-wingwall.spec.yaml")
    spec = yaml.safe_load(src.read_text())
    spec["params"]["tables_file"] = str(
        (src.parent / spec["params"]["tables_file"]).resolve())
    spec["paycodes"] = str((src.parent / spec["paycodes"]).resolve())
    spec["outputs"] = {
        "dxf": str(tmp_path / "ww.dxf"),
        "ifc": str(tmp_path / "ww.ifc"),
        "takeoff_map": str(tmp_path / "ww-takeoff-map.yaml"),
    }
    p = tmp_path / "spec.yaml"
    p.write_text(yaml.safe_dump(spec))
    assert build_from_spec(p, check_only=False) == 0
    assert (tmp_path / "ww.ifc").exists()


def test_wingwall_example_fixture_schema():
    from pathlib import Path
    import yaml
    from bimforge.author.box_section import parse_bar_spec
    from bimforge.author.wingwall_steel import (
        WingwallSteel, parse_bar_size, parse_bar_total)
    t = yaml.safe_load((Path(__file__).resolve().parents[1] / "examples"
                        / "caltrans-culvert-wingwall-tables.example.yaml").read_text())
    steel = t["steel"]
    for h, row in steel["type_e"].items():
        parse_bar_spec(row["k_bar"]); parse_bar_size(row["L_bar"])
        assert int(row["L_each_wall"]) >= 1, h
    for h, row in steel["warped"]["wall"].items():
        for sl in row["slopes"].values():
            parse_bar_spec(sl["front"]); parse_bar_spec(sl["rear"])
    for l, col in steel["warped"]["beam"]["by_l"].items():
        for reg in col:
            parse_bar_total(reg["total"])
    WingwallSteel(steel)                          # constructs cleanly


def test_wingwall_example_full_build_with_bbs(tmp_path):
    from pathlib import Path
    import yaml
    pytest.importorskip("openpyxl")           # optional, non-bundled dependency
    from openpyxl import load_workbook
    from bimforge.author.families.registry import build_from_spec
    src = (Path(__file__).resolve().parents[1]
           / "examples" / "caltrans-culvert-wingwall.spec.yaml")
    spec = yaml.safe_load(src.read_text())
    spec["params"]["tables_file"] = str(
        (src.parent / spec["params"]["tables_file"]).resolve())
    spec["paycodes"] = str((src.parent / spec["paycodes"]).resolve())
    spec["outputs"] = {"ifc": str(tmp_path / "ww.ifc"),
                       "takeoff_map": str(tmp_path / "ww-map.yaml"),
                       "bbs": str(tmp_path / "ww.bbs.xlsx")}
    p = tmp_path / "spec.yaml"
    p.write_text(yaml.safe_dump(spec))
    assert build_from_spec(p, check_only=False) == 0
    wb = load_workbook(tmp_path / "ww.bbs.xlsx")
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("Mark", "Size", "Count", "Length (ft)", "Weight (lb)")
    assert rows[-1][0] == "TOTAL" and rows[-1][4] > 0     # nonzero total steel
    # several per-mark rows (k/L/front/rear/beam/dowel/spacer across WWE+WWW)
    assert len(rows) > 4
