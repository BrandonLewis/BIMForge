"""Verifiable per-element takeoff workbook from IFC models
(`est-bim takeoff-workbook`).

One row per modelled element per pay code, with the estimator's fixed
takeoff layout — Qty, #_surfaces, L, W, H, CY, Forming SF — and then the
checks: the model's own analytic volume beside the L × W × H × Qty box
product with a variance column, the SFCA face split (top / soffit / side),
the per-row forming flags, and the pay quantity in the bid unit. CY,
Forming SF, the box volume, the variance and the pay quantity are LIVE
FORMULAS, never values, so a reviewer can change a blue input (Qty, L, W,
H, the forming flags) and watch the sheet re-total. A Summary sheet SUMIFS
the pay quantities by cost code against the official bid-item quantities.

Inputs are the same the `takeoff` verb reads (docs/ifc-takeoff-contract.md):
`Qto_BaseQuantities` in the file's declared units (scaled to SI by
`bimforge.bim.ifc_takeoff`), classification references keyed by pay code,
and the takeoff map for the pay basis per code. Surfaces come from
`bimforge.bim.ifc_mesh` — the geometry iterator's tessellation classified by
`classify_faces` — so "forming surface" is measured on the same solid the
volume is.

Units: the IFC is SI; the workbook is the estimator's feet/pounds domain.
Metres convert at the US survey foot (1200/3937 m — the foot
`bimforge.author.project.ProjectContext` emits with), kilograms at
0.45359237 kg/lb.

Output is byte-deterministic: the workbook's docProps timestamps are pinned
(as `emit_qto` pins them) and the zip members carry a fixed date, so a
rebuild of unchanged inputs is identical and diffs cleanly.

openpyxl is NOT a bimforge dependency (the `xlsx` extra): lazy import,
friendly error — the `emit_qto` convention.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable

#: metres per US survey foot — the foot the authored models are built in.
M_PER_FT = 1200.0 / 3937.0
KG_PER_LB = 0.45359237

#: pinned docProps created/modified (mirrors emit_qto.QTO_BUILD_STAMP).
WORKBOOK_BUILD_STAMP = datetime(2026, 8, 15, 12, 0, 0)
#: fixed zip-member date (the earliest DOS date); openpyxl otherwise stamps
#: every member with wall-clock time, which defeats a byte-identical rebuild.
_ZIP_DATE = (1980, 1, 1, 0, 0, 0)

FONT_NAME = "Arial"

#: The Elements sheet header, in order. The identity columns first, then the
#: estimator's takeoff layout, then the checks. This order is the contract
#: (the bim-to-takeoff skill documents it); tests pin it.
ELEMENT_HEADERS = [
    "Structure", "Element", "Description", "IFC class", "Facility part",
    "Cost code", "Item #", "Bid item", "Pay unit",
    "Qty", "#_surfaces", "L (ft)", "W (ft)", "H (ft)", "CY", "Forming SF",
    "CF (model)", "CF (L×W×H×Qty)", "Δ box vs model",
    "Area SF", "Weight LB",
    "Top SF", "Soffit SF", "Side SF", "Form sides (1/0)", "Form soffit (1/0)",
    "Pay qty", "GUID", "Model file",
    "Pay measure", "Pay measure (model)",
]

SUMMARY_HEADERS = ["Item #", "Cost code", "Bid item", "Pay unit",
                   "Modelled qty", "Official qty", "Variance", "Variance %",
                   "Elements"]

#: Forming defaults by IFC class: (form sides, form soffit). Footings, walls,
#: columns and railings form their sides; cast-in-place slabs and beams form
#: sides and soffit; precast, coverings, earthworks, pavements, rebar, proxies
#: and geographic/reference elements form nothing. A class not listed forms
#: nothing. Every row's flags are blue inputs the reviewer overrides.
FORM_DEFAULT = {
    "IfcFooting": (1, 0), "IfcWall": (1, 0), "IfcWallStandardCase": (1, 0),
    "IfcColumn": (1, 0), "IfcRailing": (1, 0),
    "IfcBeam": (1, 1), "IfcSlab": (1, 1),
    "IfcCovering": (0, 0), "IfcEarthworksCut": (0, 0),
    "IfcEarthworksFill": (0, 0), "IfcEarthworksElement": (0, 0),
    "IfcPavement": (0, 0), "IfcReinforcingBar": (0, 0),
    "IfcReinforcingElement": (0, 0), "IfcBuildingElementProxy": (0, 0),
    "IfcGeographicElement": (0, 0),
}
#: A predefined type naming precast work gets no forming whatever the class.
_PRECAST_TYPES = ("PRECAST", "PRECASTPANEL", "PRECAST_PANEL")

#: Per pay-unit factor applied to the feet/pounds-domain cell, as a formula
#: suffix. The source cell per dimension is chosen by `pay_formula`.
_UNIT_SUFFIX = {
    "volume": {"CY": "/27", "CF": "", "M3": f"*{M_PER_FT ** 3!r}",
               "CM": f"*{M_PER_FT ** 3!r}"},
    "area": {"SF": "", "SY": "/9", "M2": f"*{M_PER_FT ** 2!r}",
             "SM": f"*{M_PER_FT ** 2!r}"},
    "length": {"LF": "", "FT": "", "YD": "/3", "M": f"*{M_PER_FT!r}",
               "LM": f"*{M_PER_FT!r}"},
    "count": {"EA": "", "EACH": "", "LS": ""},
    "mass": {"LB": "", "LBS": "", "CWT": "/100", "TON": "/2000",
             "KG": f"*{KG_PER_LB!r}"},
}


@dataclass
class ElementRow:
    """One IFC element's identity and feet/pounds-domain measurements."""

    structure: str
    model_file: str
    element: str
    guid: str
    ifc_class: str
    predefined_type: str | None
    description: str
    part: str
    codes: list[str] = field(default_factory=list)
    count: float = 1.0
    length_ft: float | None = None
    width_ft: float | None = None
    height_ft: float | None = None
    volume_cf: float | None = None
    area_sf: float | None = None
    weight_lb: float | None = None
    top_sf: float | None = None
    soffit_sf: float | None = None
    side_sf: float | None = None
    surfaces: int | None = None
    quantities: dict[str, float] = field(default_factory=dict)


# ----------------------------------------------------------------------------
# Reading the model
# ----------------------------------------------------------------------------
def planar_face_count(verts, faces, *, normal_dp: int = 3,
                      offset_dp: int = 2) -> int:
    """Number of distinct planar faces in a triangle mesh: triangles grouped
    by (rounded unit normal, rounded plane offset). A box is 6 whatever its
    triangulation; a battered wall is 6; an octagonal column is 10."""
    import numpy as np
    v = np.asarray(verts, dtype=float).reshape(-1, 3)
    t = np.asarray(faces, dtype=np.int64).reshape(-1, 3)
    if len(v) == 0 or len(t) == 0:
        return 0
    v = v - v.min(axis=0)          # offsets in a local frame: no State Plane cancellation
    a, b, c = v[t[:, 0]], v[t[:, 1]], v[t[:, 2]]
    n = np.cross(b - a, c - a)
    ln = np.linalg.norm(n, axis=1)
    keep = ln > 1e-12
    n = n[keep] / ln[keep][:, None]
    d = np.einsum("ij,ij->i", n, a[keep])
    planes = {(round(float(x), normal_dp), round(float(y), normal_dp),
               round(float(z), normal_dp), round(float(o), offset_dp))
              for (x, y, z), o in zip(n, d)}
    return len(planes)


def surfaces_of(ifc_path: str, *, skip_classes=("IfcGeographicElement",),
                flat: float = 0.85) -> dict[str, dict]:
    """{GlobalId: {top_sf, soffit_sf, side_sf, surfaces}} in FEET for every
    element the geometry iterator tessellates (reference terrain skipped)."""
    from bimforge.bim.ifc_mesh import classify_faces, element_meshes

    out = {}
    for guid, _name, _cls, verts, faces in element_meshes(
            str(ifc_path), skip_classes=skip_classes):
        v_ft = verts / M_PER_FT
        cls = classify_faces(v_ft, faces, flat=flat)
        out[guid] = {"top_sf": cls["top_sf"], "soffit_sf": cls["bottom_sf"],
                     "side_sf": cls["side_sf"],
                     "surfaces": planar_face_count(v_ft, faces)}
    return out


def _container_of(product):
    for rel in getattr(product, "ContainedInStructure", None) or []:
        return rel.RelatingStructure
    return None


def _structure_name(product) -> str | None:
    """The nearest IfcBuilding / IfcFacility (bridge, road, ...) above the
    element, by name — the structure emit_combined_ifc names each building
    after. None when the spatial tree has no such ancestor."""
    node = _container_of(product)
    seen = 0
    while node is not None and seen < 32:
        if node.is_a("IfcBuilding") or node.is_a("IfcFacility"):
            return node.Name or None
        parents = getattr(node, "Decomposes", None) or []
        node = parents[0].RelatingObject if parents else None
        seen += 1
    return None


def _weight_kg(qtys: dict) -> float | None:
    for name in ("GrossWeight", "NetWeight", "Weight"):
        if name in qtys:
            return qtys[name]
    return None


def read_elements(ifc_path: str, *, model_label: str | None = None
                  ) -> list[ElementRow]:
    """Every IfcElement with quantities or geometry, measured into the
    feet/pounds domain, with its SFCA surfaces from the tessellation."""
    import ifcopenshell

    from bimforge.bim.ifc_takeoff import (_unit_scales, classifications_of,
                                          quantities_of)

    path = str(ifc_path)
    model = ifcopenshell.open(path)
    scales = _unit_scales(model)
    faces = surfaces_of(path)
    stem = Path(path).stem
    rows = []
    for el in model.by_type("IfcElement"):
        qtys = quantities_of(el, scales)
        if not qtys and not el.Representation:
            continue
        fx = faces.get(el.GlobalId, {})
        container = _container_of(el)
        part = (container.Name or container.is_a()) if container is not None else ""
        kg = _weight_kg(qtys)

        def _ft(name, power=1):
            v = qtys.get(name)
            return None if v is None else v / M_PER_FT ** power

        rows.append(ElementRow(
            structure=model_label or _structure_name(el) or stem,
            model_file=path,
            element=el.Name or f"{el.is_a()}-{el.GlobalId[:6]}",
            guid=el.GlobalId, ifc_class=el.is_a(),
            predefined_type=getattr(el, "PredefinedType", None),
            description=el.Description or "", part=part,
            codes=sorted({c["code"] for c in classifications_of(el) if c["code"]}),
            count=float(qtys.get("Count", 1.0)),
            length_ft=_ft("Length"), width_ft=_ft("Width"),
            height_ft=_ft("Height"), volume_cf=_ft("GrossVolume", 3),
            area_sf=_ft("GrossArea", 2),
            weight_lb=None if kg is None else kg / KG_PER_LB,
            top_sf=fx.get("top_sf"), soffit_sf=fx.get("soffit_sf"),
            side_sf=fx.get("side_sf"), surfaces=fx.get("surfaces"),
            quantities=qtys))
    # Entity-id order is an emitter accident; structure / facility part /
    # name is how an estimator reads the sheet, and it is stable.
    rows.sort(key=lambda e: (e.structure, e.part, e.element, e.guid))
    return rows


def forming_default(ifc_class: str, predefined_type: str | None = None
                    ) -> tuple[int, int]:
    """(form sides, form soffit) for an element, by IFC class — precast
    predefined types form nothing."""
    if predefined_type and str(predefined_type).upper() in _PRECAST_TYPES:
        return (0, 0)
    return FORM_DEFAULT.get(ifc_class, (0, 0))


# ----------------------------------------------------------------------------
# The takeoff map and the official quantities
# ----------------------------------------------------------------------------
def load_map(path: str | None) -> dict:
    """The takeoff map document (validated by `ifc_takeoff.normalize_map`)."""
    import yaml

    if not path:
        return {}
    with open(path, encoding="utf-8") as fh:
        return yaml.safe_load(fh) or {}


def code_entries(tmap: dict) -> dict[str, dict]:
    """{code: {measure, unit, item_number, description}} — the pay basis per
    code. A code with plural `measures:` contributes its FIRST entry (one pay
    line per element row); the map's `defaults:` fill missing keys."""
    from bimforge.bim.ifc_takeoff import normalize_map

    codes_cfg, defaults, entries_by_code = normalize_map(tmap)
    out = {}
    for code, cfg in codes_cfg.items():
        cfg = dict(cfg or {})
        entries = entries_by_code.get(code)
        if entries:
            first = dict(entries[0])
            cfg.pop("measures", None)
            cfg.update({k: v for k, v in first.items() if v is not None})
        merged = dict(defaults)
        merged.update(cfg)
        out[str(code)] = merged
    return out


def load_official(path: str | None) -> dict[str, dict]:
    """{item_number: {quantity, description, unit}} from a JSON bid-item
    list — `[{item_number, quantity | estimated_quantity, description?,
    unit?}]`, or an object holding that list under `bid_items`, `items` or
    `bid_frame`. Missing file argument → {} (Official qty stays blank)."""
    if not path:
        return {}
    with open(path, encoding="utf-8") as fh:
        doc = json.load(fh)
    items = doc
    if isinstance(doc, dict):
        for key in ("bid_items", "items", "bid_frame", "rows"):
            if isinstance(doc.get(key), list):
                items = doc[key]
                break
        else:
            raise ValueError(f"{path}: expected a list of bid items (or an "
                             f"object with a bid_items/items/bid_frame list)")
    if not isinstance(items, list):
        raise ValueError(f"{path}: expected a list of bid items")
    out = {}
    for it in items:
        if not isinstance(it, dict) or it.get("item_number") in (None, ""):
            continue
        qty = it.get("quantity")
        if qty is None:
            qty = it.get("estimated_quantity")
        try:
            qty = None if qty in (None, "") else float(qty)
        except (TypeError, ValueError):
            raise ValueError(f"{path}: item {it['item_number']}: quantity "
                             f"{qty!r} is not a number") from None
        out[str(it["item_number"])] = {
            "quantity": qty, "description": it.get("description") or "",
            "unit": it.get("unit") or it.get("pay_unit") or ""}
    return out


def pay_formula(cfg: dict, row: int, col: dict[str, str]) -> str:
    """The Pay qty cell formula for one element row: the feet/pounds-domain
    source cell for the code's measure, scaled to its pay unit. '' when the
    code has no measure (an unmapped code pays nothing — review it)."""
    from bimforge.bim.ifc_takeoff import MEASURE_DIM

    measure = cfg.get("measure")
    dim = MEASURE_DIM.get(measure or "")
    if dim is None:
        return ""
    unit = str(cfg.get("unit") or "").upper()
    suffix = _UNIT_SUFFIX[dim].get(unit)
    if suffix is None:
        raise ValueError(f"pay unit {unit or '(none)'!r} is not a {dim} unit "
                         f"(choices: {sorted(_UNIT_SUFFIX[dim])})")
    if dim == "volume":
        if measure != "GrossVolume":
            return f"={col['Pay measure (model)']}{row}{suffix}"
        src = f"{col['CY']}{row}" if unit == "CY" else f"{col['CF (model)']}{row}"
        return f"={src}" if unit == "CY" else f"={src}{suffix}"
    if dim == "area":
        key = "Area SF" if measure == "GrossArea" else "Pay measure (model)"
        return f"={col[key]}{row}{suffix}"
    if dim == "length":
        key = {"Length": "L (ft)", "Width": "W (ft)", "Height": "H (ft)"}.get(
            measure, "Pay measure (model)")
        return f"={col[key]}{row}*{col['Qty']}{row}{suffix}"
    if dim == "count":
        if unit == "LS" or cfg.get("count_as") is not None:
            value = cfg.get("count_as", 1.0)
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError("count_as must be a number")
            code = col["Cost code"]
            # Count overrides and lump sums apply once to the code, rather
            # than once for every solid that backs that pay item.
            return f'=IF(COUNTIF(${code}$2:{code}{row},{code}{row})=1,{float(value)!r},0)'
        return f"={col['Qty']}{row}"
    return f"={col['Pay measure (model)']}{row}{suffix}"


def pay_measure_value(row: ElementRow, cfg: dict) -> float | None:
    """Keep the requested Qto measure distinct from others of its dimension.

    The extra audit cell uses the same feet/pounds domain as the main
    columns. Never replace a missing Net quantity with its Gross value.
    """
    from bimforge.bim.ifc_takeoff import MEASURE_DIM, resolve_measure

    measure = cfg.get("measure")
    dim = MEASURE_DIM.get(measure or "")
    if dim is None:
        return None
    actual = resolve_measure(row.quantities, measure)
    if actual is None:
        raise ValueError(f"element {row.element!r}: lacks mapped measure {measure}")
    value = row.quantities[actual]
    if dim in ("length", "area", "volume"):
        return value / M_PER_FT ** {"length": 1, "area": 2, "volume": 3}[dim]
    return value / KG_PER_LB if dim == "mass" else value


# ----------------------------------------------------------------------------
# Writing the workbook
# ----------------------------------------------------------------------------
def _save_deterministic(wb, path: str) -> None:
    """Save with pinned docProps AND fixed zip-member dates: openpyxl stamps
    each member with wall-clock time, and its save overwrites
    ``properties.modified`` with now, so ``docProps/core.xml``'s modified
    date is rewritten here too; otherwise two rebuilds a second apart are
    byte-different."""
    wb.properties.created = wb.properties.modified = WORKBOOK_BUILD_STAMP
    stamp = WORKBOOK_BUILD_STAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=_ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + stamp + rb"\g<2>", data)
            dst.writestr(zi, data)


def _item_sort_key(item: str):
    return (0, int(item)) if item.isdigit() else (1, item)


def build_workbook(ifc_paths: Iterable[str], tmap: dict, out: str, *,
                   official: dict[str, dict] | None = None,
                   log=None) -> dict:
    """Read every IFC, write the Elements / Summary / Legend / Sources
    workbook to `out`, and return {"elements": rows, "summary": rows,
    "models": n}. `official` is `load_official`'s dict (Official qty blank
    without it); `log` receives review notes (default: stderr)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError(
            "takeoff-workbook needs openpyxl, which bimforge does not "
            "bundle: pip install openpyxl") from None

    log = log or (lambda msg: print(msg, file=sys.stderr))
    official = official or {}
    codes_cfg = code_entries(tmap)
    ifc_paths = [str(p) for p in ifc_paths]
    if not ifc_paths:
        raise ValueError("takeoff-workbook: no IFC files given")

    rows: list[ElementRow] = []
    for p in ifc_paths:
        if not Path(p).is_file():
            raise OSError(f"{p}: no such IFC file")
        rows.extend(read_elements(p))

    blue = Font(name=FONT_NAME, color="0000FF")
    black = Font(name=FONT_NAME)
    bold = Font(name=FONT_NAME, bold=True)
    head_fill = PatternFill("solid", fgColor="DDDDDD")

    def _font_all(ws):
        for row_ in ws.iter_rows():
            for cell in row_:
                if cell.font is None or cell.font.name != FONT_NAME:
                    cell.font = black

    wb = Workbook()
    # ---------------------------------------------------------------- Elements
    ws = wb.active
    ws.title = "Elements"
    ws.append(ELEMENT_HEADERS)
    for c in range(1, len(ELEMENT_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    col = {h: get_column_letter(i + 1) for i, h in enumerate(ELEMENT_HEADERS)}

    unmapped: set[str] = set()
    r = 1
    for row in rows:
        for code in (row.codes or [""]):       # one line per pay code
            r += 1
            cfg = codes_cfg.get(code, {})
            if code and not cfg:
                unmapped.add(code)
            item = str(cfg.get("item_number") or "")
            off = official.get(item, {})
            fs, fb = forming_default(row.ifc_class, row.predefined_type)
            ws.append([
                row.structure, row.element, row.description, row.ifc_class,
                row.part, code or "(reference)", item,
                off.get("description") or cfg.get("description") or "",
                cfg.get("unit") or "",
                row.count, row.surfaces, row.length_ft, row.width_ft,
                row.height_ft, None, None,
                row.volume_cf, None, None, row.area_sf, row.weight_lb,
                row.top_sf, row.soffit_sf, row.side_sf, fs, fb, None,
                row.guid, row.model_file,
                cfg.get("measure") or "", pay_measure_value(row, cfg)])
            L, W, H, Q = (f"{col[k]}{r}" for k in ("L (ft)", "W (ft)", "H (ft)", "Qty"))
            cf_model, cf_box = f"{col['CF (model)']}{r}", f"{col['CF (L×W×H×Qty)']}{r}"
            side, soffit = f"{col['Side SF']}{r}", f"{col['Soffit SF']}{r}"
            f_sides, f_soffit = f"{col['Form sides (1/0)']}{r}", f"{col['Form soffit (1/0)']}{r}"
            ws[cf_box] = (f'=IF(AND(ISNUMBER({L}),ISNUMBER({W}),ISNUMBER({H})),'
                          f'{L}*{W}*{H}*{Q},"")')
            ws[f"{col['Δ box vs model']}{r}"] = (
                f'=IF(AND(ISNUMBER({cf_box}),ISNUMBER({cf_model}),{cf_model}<>0),'
                f'{cf_box}/{cf_model}-1,"")')
            ws[f"{col['CY']}{r}"] = f'=IF(ISNUMBER({cf_model}),{cf_model}/27,"")'
            ws[f"{col['Forming SF']}{r}"] = (
                f'=IF(ISNUMBER({side}),{side}*{f_sides}+{soffit}*{f_soffit},"")')
            ws[f"{col['Pay qty']}{r}"] = pay_formula(cfg, r, col) if code else ""
            for k in ("Qty", "L (ft)", "W (ft)", "H (ft)", "Form sides (1/0)",
                      "Form soffit (1/0)"):
                ws[f"{col[k]}{r}"].font = blue
            for k in ("CY", "Forming SF", "CF (model)", "CF (L×W×H×Qty)",
                      "Area SF", "Weight LB", "Top SF", "Soffit SF", "Side SF",
                      "Pay qty", "Pay measure (model)"):
                ws[f"{col[k]}{r}"].number_format = "#,##0.00"
            ws[f"{col['Δ box vs model']}{r}"].number_format = "0.0%"
            for k in ("L (ft)", "W (ft)", "H (ft)"):
                ws[f"{col[k]}{r}"].number_format = "0.000"
            ws[f"{col['Qty']}{r}"].number_format = "0"
    last = max(r, 2)
    for c in range(1, len(ELEMENT_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions[col["Element"]].width = 34
    ws.column_dimensions[col["Description"]].width = 30
    ws.column_dimensions[col["Bid item"]].width = 36
    ws.column_dimensions[col["GUID"]].width = 24
    ws.freeze_panes = "C2"
    _font_all(ws)

    # ---------------------------------------------------------------- Summary
    sm = wb.create_sheet("Summary")
    sm.append(SUMMARY_HEADERS)
    for c in range(1, len(SUMMARY_HEADERS) + 1):
        sm.cell(row=1, column=c).font = bold
        sm.cell(row=1, column=c).fill = head_fill
    rng_code = f"Elements!${col['Cost code']}$2:${col['Cost code']}${last}"
    rng_pay = f"Elements!${col['Pay qty']}$2:${col['Pay qty']}${last}"
    all_codes = set(codes_cfg) | {c for row in rows for c in row.codes}
    ordered = sorted(all_codes, key=lambda c: (
        _item_sort_key(str(codes_cfg.get(c, {}).get("item_number") or "")), c))
    rr = 1
    for code in ordered:
        rr += 1
        cfg = codes_cfg.get(code, {})
        item = str(cfg.get("item_number") or "")
        off = official.get(item, {})
        sm.append([item, code,
                   off.get("description") or cfg.get("description") or "",
                   cfg.get("unit") or "", None, off.get("quantity"),
                   None, None, None])
        sm[f"E{rr}"] = f"=SUMIFS({rng_pay},{rng_code},B{rr})"
        sm[f"G{rr}"] = f'=IF(ISNUMBER(F{rr}),E{rr}-F{rr},"")'
        sm[f"H{rr}"] = f'=IF(AND(ISNUMBER(F{rr}),F{rr}<>0),G{rr}/F{rr},"")'
        sm[f"I{rr}"] = f"=COUNTIFS({rng_code},B{rr})"
        sm[f"F{rr}"].font = blue
        for c in "EFG":
            sm[f"{c}{rr}"].number_format = "#,##0.00"
        sm[f"H{rr}"].number_format = "0.0%"
    for c, w in zip("ABCDEFGHI", (8, 12, 44, 9, 14, 14, 12, 11, 10)):
        sm.column_dimensions[c].width = w
    sm.freeze_panes = "A2"
    _font_all(sm)

    # ---------------------------------------------------------------- Legend
    lg = wb.create_sheet("Legend")
    lines = [
        "Takeoff workbook generated from the IFC models listed on Sources; "
        "every modelled element is one row per pay code on Elements.",
        "Blue cells are inputs you may edit (Qty, L, W, H, Form sides, Form "
        "soffit, Official qty); black cells are formulas or model values.",
        "CF (L×W×H×Qty) is the hand check; CF (model) is the analytic volume "
        "of the emitted solid (Qto_BaseQuantities GrossVolume). Δ box vs "
        "model shows when the solid is not a box (battered, skewed, voided).",
        "#_surfaces is the number of distinct planar faces of the element's "
        "tessellation. Top / Soffit / Side SF come from bimforge's SFCA face "
        "classifier on the same solid (|nz| > 0.85 = top or soffit, else side, "
        "measured on the true slope).",
        "Forming SF = Side SF × Form sides + Soffit SF × Form soffit. Defaults "
        "by IFC class: footings, walls, columns, railings form sides; "
        "cast-in-place slabs and beams form sides + soffit; precast, coverings, "
        "earthworks, pavements, rebar, proxies and geographic elements form "
        "nothing. Override per row.",
        "Pay qty converts the model quantity to the bid unit per the takeoff "
        "map's measure and unit for the cost code (CY = CF/27, SY = SF/9, "
        "TON = LB/2000, LF = L × Qty, EA = Qty). A cost code with no map "
        "entry pays nothing and is listed on Summary with a blank unit.",
        "Pay measure names the requested Qto quantity. Pay measure (model) "
        "is that exact value in feet, square feet, cubic feet, pounds or "
        "each, preserving net/gross and length/width/height distinctions.",
        "Summary sums Pay qty by cost code (SUMIFS) and compares it to the "
        "official bid-item quantity. A large variance is a finding to resolve "
        "against the plans, not a number to tune.",
        "Units: IFC metres → US survey feet (1200/3937 m); kilograms → pounds "
        "(0.45359237 kg/lb).",
    ]
    for i, text in enumerate(lines, 1):
        lg[f"A{i}"] = text
    lg.column_dimensions["A"].width = 140
    _font_all(lg)

    # ---------------------------------------------------------------- Sources
    src = wb.create_sheet("Sources")
    src.append(["Model", "File", "SHA-256"])
    for c in range(1, 4):
        src.cell(row=1, column=c).font = bold
        src.cell(row=1, column=c).fill = head_fill
    for p in ifc_paths:
        src.append([Path(p).stem, p,
                    hashlib.sha256(Path(p).read_bytes()).hexdigest()])
    for c, w in zip("ABC", (28, 60, 68)):
        src.column_dimensions[c].width = w
    _font_all(src)

    _save_deterministic(wb, out)
    for code in sorted(unmapped):
        log(f"review: cost code {code!r} has no takeoff-map entry — its rows "
            f"pay nothing until the map names a measure and unit")
    return {"elements": last - 1 if rows else 0, "summary": rr - 1,
            "models": len(ifc_paths)}
